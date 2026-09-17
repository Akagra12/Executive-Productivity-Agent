import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(15, 23, 42)        # #0f172a Deep Slate
    CARD_BG = RGBColor(30, 41, 59)         # #1e293b Card Slate
    CARD_BORDER = RGBColor(51, 65, 85)     # #334155
    ACCENT_INDIGO = RGBColor(129, 140, 248)# #818cf8 Accent Indigo
    ACCENT_CYAN = RGBColor(56, 189, 248)   # #38bdf8 Accent Cyan
    ACCENT_GREEN = RGBColor(52, 211, 153)  # #34d399 Accent Green
    ACCENT_AMBER = RGBColor(251, 191, 36)  # #fbbf24 Accent Amber
    ACCENT_RED = RGBColor(248, 113, 113)   # #f87171 Accent Red
    TEXT_WHITE = RGBColor(248, 250, 252)   # #f8fafc Primary Text
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94a3b8 Secondary Text

    blank_layout = prs.slide_layouts[6]

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title_text, category_text="AIONOS EXECUTIVE PRODUCTIVITY AGENT"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_INDIGO

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, title, items, badge_color=ACCENT_INDIGO, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = badge_color
        p0.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_WHITE
            p.space_after = Pt(6)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 1: Title Slide
    # ═══════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "AIONOS AGENTIC AI FACTORY"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_INDIGO
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Executive Productivity Agent"
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Multi-Modal Ingestion · Deterministic Deadlines · Grounded Daily Action Briefs"
    p3.font.size = Pt(18)
    p3.font.color.rgb = ACCENT_CYAN
    p3.space_after = Pt(24)

    p4 = tf1.add_paragraph()
    p4.text = "Presenter: Akagra  |  Target Executive: Arjun Malhotra, VP Product & Strategy\nGitHub: https://github.com/Akagra12/Executive-Productivity-Agent"
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_MUTED

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 2: Problem Statement
    # ═══════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, "The Executive Information Overload Challenge")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Fragmented Data Silos",
             ["Action items scattered across transcripts, emails, calendars, voice notes",
              "No unified view of promises made vs promises received",
              "Context lost across different messaging formats"], ACCENT_AMBER)

    add_card(s2, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Missed Time-of-Day Deadlines",
             ["Traditional date trackers miss 09:00 AM morning deadlines",
              "Deliverables get lost in email chains",
              "Lack of real-time proactive escalation"], ACCENT_RED)

    add_card(s2, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.4),
             "Ambiguous Task Ownership",
             ["Meeting discussions frequently end without clear owners",
              "Standard AI models invent fake owners or hallucinate facts",
              "Critical agreements fall through corporate cracks"], ACCENT_RED)

    add_card(s2, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.4),
             "Lack of Source Auditability",
             ["Executives cannot verify if AI summaries are accurate",
              "Zero quote citations or turn traceability in legacy summaries",
              "High risk of misinformation in high-stakes decisions"], ACCENT_AMBER)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 3: Solution Overview
    # ═══════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, "The AIONOS Executive Intelligence Layer")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "1. Multi-Modal Ingestion",
             ["Ingests 71 raw inputs",
              "Transcripts, emails, calendars & audio voice notes",
              "Normalized SourceMessage schema",
              "100% source fidelity"], ACCENT_CYAN)

    add_card(s3, Inches(3.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "2. Deadline Engine",
             ["Time-of-day precision (09:00 vs 18:00)",
              "Concluded meetings marked past_event",
              "5-Day workweek simulation",
              "Dynamic overdue badges"], ACCENT_GREEN)

    add_card(s3, Inches(6.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "3. Ownership Tri-Partition",
             ["My Actions (Arjun)",
              "Waiting on Others (Team)",
              "Unclear Ownership (Mumbai lease)",
              "Zero hallucinated owners"], ACCENT_AMBER)

    add_card(s3, Inches(9.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "4. Grounded AI Q&A",
             ["Context-constrained RAG",
              "Verbatim quote citations",
              "Exact speaker turn links",
              "Zero hallucinations"], ACCENT_INDIGO)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 4: Architecture
    # ═══════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, "End-to-End Modular System Architecture")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.5),
             "Executive Frontend UI (React 19 + Vite)",
             ["Daily Action Brief  |  Commitment Matrix  |  Raw Data Explorer  |  Deduplication Inspector  |  AI Query Assistant  |  5-Day Time Switcher"], ACCENT_INDIGO)

    add_card(s4, Inches(0.8), Inches(3.5), Inches(5.7), Inches(3.3),
             "Core Intelligence & Rule Engines",
             ["Deterministic Deadline Engine (Time-of-day calculations)",
              "Multi-Signal Deduplication Engine (Jaccard lexical scoring)",
              "Executive Brief Synthesizer (Priority scoring: 95 > 85 > 75)",
              "Grounded LLM & Retrieval Engine (Zero-hallucination Q&A)"], ACCENT_CYAN)

    add_card(s4, Inches(6.8), Inches(3.5), Inches(5.7), Inches(3.3),
             "Multi-Modal Ingestion & Normalizer",
             ["Meeting Transcripts: Speaker turns & timestamp mapping",
              "Email Threads: Headers, threads, sequence tracking",
              "Calendars: Attendees, start/end times & event types",
              "Voice Notes: Personal dictation transcripts & memos"], ACCENT_GREEN)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 5: Ingestion & Raw Data Pack
    # ═══════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, "Multi-Modal Ingestion & 100% Traceability")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Meeting Transcripts (8 Turns)",
             ["Monday Leadership Sync (2026-09-21)",
              "Mapped speaker dialogue to individual turn IDs",
              "Captures informal verbal commitments"], ACCENT_CYAN)

    add_card(s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Email Threads (5 Threads)",
             ["Multi-party correspondence with Raghav, Neha, external partners",
              "Preserves email headers, timestamps, and thread context",
              "Tracks deliverable handoffs"], ACCENT_INDIGO)

    add_card(s5, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.4),
             "Calendar Records (4 Calendars)",
             ["Cross-references schedules for Arjun, Raghav, Neha, Divya",
              "Differentiates external calls, 1:1s, and focus blocks",
              "Evaluates attendance vs deliverable deadlines"], ACCENT_GREEN)

    add_card(s5, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.4),
             "Voice Notes (2 Memos)",
             ["Dictated memos recorded while commuting",
              "Personal reminders with high strategic importance",
              "Consolidated with email threads via deduplication"], ACCENT_AMBER)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 6: Deadline Engine & Time Simulation
    # ═══════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, "Deterministic Deadline & Time-of-Day Intelligence")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Time-of-Day Precision",
             ["Evaluates exact HH:MM deadlines",
              "Vendor list (c_001) due at 09:00 AM",
              "At 10:00 AM on Sept 23 -> Evaluates as Overdue",
              "EOD deadlines (c_002) valid until 18:00"], ACCENT_RED)

    add_card(s6, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Calendar Attendance Logic",
             ["Treats meetings separately from tasks",
              "Board Prep (c_007) scheduled Sept 24 09:00-10:00",
              "Before meeting -> Upcoming / Due Today",
              "After meeting concludes -> Clean past_event (never overdue)"], ACCENT_GREEN)

    add_card(s6, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "5-Day Time Simulation",
             ["Mon 21 Sep -> 0 Overdue, Sync meetings",
              "Wed 23 Sep -> Vendor list Overdue",
              "Thu 24 Sep -> Board Prep morning",
              "Fri 25 Sep -> Mumbai Lease deadline"], ACCENT_CYAN)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 7: Ownership Partition & Deduplication
    # ═══════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    add_header(s7, "Strict Ownership Partition & Deduplication")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "My Actions vs Waiting on Others",
             ["Direct owner items assigned to Arjun (budget model, vendor review)",
              "Delegated items waiting on Raghav (pitch deck) and Divya (compliance)",
              "Fast API filtering by category, person, and role"], ACCENT_INDIGO)

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Unclear Ownership Protection",
             ["Mumbai Office Lease (c_004) mentioned without owner",
              "System strictly sets owner: null and ownership_unclear: true",
              "Calculates deadline urgency without guessing an owner"], ACCENT_AMBER)

    add_card(s7, Inches(0.8), Inches(4.4), Inches(11.7), Inches(2.4),
             "Multi-Signal Deduplication Inspector",
             ["Cross-references extracted items across emails, meetings, and voice memos",
              "Combines owner match (30%), recipient match (20%), deadline match (20%), and text similarity (30%)",
              "Consolidates duplicate promises into single canonical commitments with merged evidence trails"], ACCENT_GREEN)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 8: Daily Executive Brief Synthesizer
    # ═══════════════════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    add_header(s8, "Executive Daily Brief Synthesizer")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "Priority Ordering Hierarchy",
             ["1. Overdue Items (Score 95) -> Immediate attention",
              "2. Due Today (Score 85) -> Priority execution",
              "3. Unclear Ownership (Score 75) -> High-risk blockers",
              "4. Scheduled Meetings (Score 65) -> Calendar syncs",
              "5. Upcoming Deliverables (Score 55) -> 1-3 day horizon",
              "6. Waiting on Others (Score 45) -> Follow-up list",
              "7. Completed Commitments (Score 15) -> Resolved archive"], ACCENT_AMBER)

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "Multi-Format Executive Delivery",
             ["Interactive Web Dashboard -> Real-time status badges",
              "Rich Markdown Export -> Document summary for executive notebooks",
              "WhatsApp / Slack Chat Snippet -> Compact text for mobile review",
              "Structured JSON -> Machine-readable payload for enterprise workflow integrations",
              "One-click clipboard export with zero duplicate lines"], ACCENT_CYAN)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 9: Grounded Assistant & Verification
    # ═══════════════════════════════════════════════════════════════════════════
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)
    add_header(s9, "Grounded AI Assistant & 100% Test Coverage")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "Zero-Hallucination Natural Language Q&A",
             ["'What did I promise Raghav?'\n  -> Series B narrative feedback by Tuesday EOD\n  -> Cites Turn #4 in Monday Leadership Sync",
              "'What needs action today?'\n  -> Lists active deliverables based on simulated date",
              "'Which tasks have unclear ownership?'\n  -> Identifies Mumbai lease without inventing owners",
              "Context-constrained reasoning guarantees truthfulness"], ACCENT_INDIGO)

    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "Automated Test Verification",
             ["test_deadlines.js: 82 / 82 PASSED (100%)\n  • Morning time-of-day accuracy\n  • Calendar past_event handling\n  • Unclear urgency countdowns",
              "test_filters.js: 27 / 27 PASSED (100%)\n  • Category, person, and role combinations",
              "test_brief.js: 5 / 5 Dates Generated with 0 duplicates",
              "Frontend Build: 1888 modules transformed with 0 errors"], ACCENT_GREEN)

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 10: Conclusion & Impact
    # ═══════════════════════════════════════════════════════════════════════════
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_background(s10)
    add_header(s10, "Summary & Reviewer Delivery")

    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "Executive Value Delivered",
             ["Saves 1-2 hours daily by automating multi-channel tracking",
              "Eliminates dropped commitments with time-of-day deadline monitoring",
              "Prevents execution blindspots by surfacing unassigned tasks",
              "Delivers audit-grade confidence with 100% verbatim source evidence"], ACCENT_CYAN)

    add_card(s10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "Prototype Access & Repository",
             ["GitHub Repository:\n  https://github.com/Akagra12/Executive-Productivity-Agent",
              "One-Command Local Run:\n  Backend: cd server && npm start (Port 3001)\n  Frontend: cd Client && npm run dev (Port 5173)",
              "Pre-packaged sample data pack and regression test suites",
              "Thank you! Ready for Live Demonstration."], ACCENT_GREEN)

    output_path = os.path.join(os.path.dirname(__file__), "AIONOS_Executive_Productivity_Agent.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
